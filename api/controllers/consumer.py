import hashlib
import os
import smtplib
import io
from datetime import date, timedelta
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from typing import Optional

from fastapi import HTTPException, UploadFile
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

SMTP_HOST = os.getenv("SMTP_HOST", "smtp.hostinger.com")
SMTP_PORT = int(os.getenv("SMTP_PORT", 465))
SMTP_USER = os.getenv("SMTP_USER", "")
SMTP_PASS = os.getenv("SMTP_PASS", "")
NOTIFY_EMAIL = os.getenv("CONSUMER_NOTIFY_EMAIL", "amit@enertsol.com")
SMTP_FROM = f"Meter Portal <{SMTP_USER}>"


def md5(password: str) -> str:
    return hashlib.md5(password.encode()).hexdigest()


def _business_days_from_today(n: int) -> str:
    d = date.today()
    added = 0
    while added < n:
        d += timedelta(days=1)
        if d.weekday() < 5:
            added += 1
    return d.strftime("%m/%d/%Y")


def _send_email(subject: str, html_body: str) -> None:
    try:
        msg = MIMEMultipart("alternative")
        msg["Subject"] = subject
        msg["From"] = SMTP_FROM
        msg["To"] = NOTIFY_EMAIL
        msg.attach(MIMEText(html_body, "html"))
        with smtplib.SMTP_SSL(SMTP_HOST, SMTP_PORT) as server:
            server.login(SMTP_USER, SMTP_PASS)
            server.sendmail(SMTP_USER, [NOTIFY_EMAIL], msg.as_string())
    except Exception:
        pass  # email failures must not break the request


# ---------------------------------------------------------------------------
# Auth
# ---------------------------------------------------------------------------

async def login_consumer(db: AsyncSession, login: str, password: str) -> dict:
    hashed = md5(password)
    result = await db.execute(
        text("""
            SELECT uid, name, email, role, status
            FROM users
            WHERE (name = :login OR email = :login) AND password = :password
            LIMIT 1
        """),
        {"login": login, "password": hashed},
    )
    user = result.fetchone()
    if not user:
        raise HTTPException(status_code=401, detail="Invalid login or password")
    if user.status == 0:
        raise HTTPException(status_code=403, detail="Account is inactive")
    return {
        "uid": user.uid,
        "name": user.name,
        "email": user.email,
        "role": user.role,
    }


# ---------------------------------------------------------------------------
# Admin — users
# ---------------------------------------------------------------------------

async def list_users(db: AsyncSession) -> list:
    result = await db.execute(
        text("""
            SELECT u.uid, u.name, u.email, u.role, u.status,
                   COUNT(ui.sr) AS meter_count
            FROM users u
            LEFT JOIN user_information ui ON u.uid = ui.uid AND ui.status != 4
            GROUP BY u.uid
            ORDER BY u.name ASC
        """)
    )
    rows = result.fetchall()
    return [
        {
            "uid": r.uid,
            "name": r.name,
            "email": r.email,
            "role": r.role,
            "status": r.status,
            "meter_count": r.meter_count,
        }
        for r in rows
    ]


async def create_user(
    db: AsyncSession, name: str, password: str, email: str
) -> dict:
    # duplicate check
    dup = await db.execute(
        text("SELECT uid FROM users WHERE name = :name OR email = :email LIMIT 1"),
        {"name": name, "email": email},
    )
    if dup.fetchone():
        raise HTTPException(status_code=400, detail="Username or email already exists")

    await db.execute(
        text("""
            INSERT INTO users (name, email, password, role, status)
            VALUES (:name, :email, :password, 2, 1)
        """),
        {"name": name, "email": email, "password": md5(password)},
    )
    await db.commit()

    result = await db.execute(
        text("SELECT uid FROM users WHERE name = :name LIMIT 1"), {"name": name}
    )
    uid = result.scalar()

    _send_email(
        subject=f"New Consumer Portal Account: {name}",
        html_body=f"""
        <p>A new account has been created on the Multi Meter Management Portal.</p>
        <table>
          <tr><td><b>Username:</b></td><td>{name}</td></tr>
          <tr><td><b>Password:</b></td><td>{password}</td></tr>
          <tr><td><b>Email:</b></td><td>{email}</td></tr>
        </table>
        """,
    )
    return {"uid": uid, "name": name, "email": email, "role": 2, "status": 1}


async def update_user(
    db: AsyncSession, uid: int, name: str, password: Optional[str], email: str
) -> None:
    if password:
        await db.execute(
            text("""
                UPDATE users SET name = :name, email = :email, password = :password
                WHERE uid = :uid
            """),
            {"name": name, "email": email, "password": md5(password), "uid": uid},
        )
    else:
        await db.execute(
            text("UPDATE users SET name = :name, email = :email WHERE uid = :uid"),
            {"name": name, "email": email, "uid": uid},
        )
    await db.commit()


async def toggle_user_status(db: AsyncSession, uid: int) -> int:
    result = await db.execute(
        text("SELECT status FROM users WHERE uid = :uid"), {"uid": uid}
    )
    row = result.fetchone()
    if not row:
        raise HTTPException(status_code=404, detail="User not found")
    new_status = 0 if row.status == 1 else 1
    await db.execute(
        text("UPDATE users SET status = :status WHERE uid = :uid"),
        {"status": new_status, "uid": uid},
    )
    await db.commit()
    return new_status


# ---------------------------------------------------------------------------
# Admin — Excel upload
# ---------------------------------------------------------------------------

async def upload_meters_excel(
    db: AsyncSession, uid: int, file: UploadFile
) -> int:
    try:
        import openpyxl
    except ImportError:
        raise HTTPException(status_code=500, detail="openpyxl not installed")

    content = await file.read()
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active

    rows = list(ws.iter_rows(min_row=2, values_only=True))
    wb.close()

    # Verify user exists
    check = await db.execute(
        text("SELECT uid FROM users WHERE uid = :uid"), {"uid": uid}
    )
    if not check.fetchone():
        raise HTTPException(status_code=404, detail="User not found")

    # Delete existing meters for this user
    await db.execute(
        text("DELETE FROM user_information WHERE uid = :uid"), {"uid": uid}
    )

    count = 0
    for row in rows:
        if not row or not row[0]:
            continue
        esid = str(row[0]).strip() if row[0] else ""
        address = str(row[1]).strip() if len(row) > 1 and row[1] else ""
        unit_number = str(row[2]).strip() if len(row) > 2 and row[2] else ""
        city = str(row[3]).strip() if len(row) > 3 and row[3] else ""
        zip_code = str(row[4]).strip() if len(row) > 4 and row[4] else ""

        if not esid:
            continue

        await db.execute(
            text("""
                INSERT INTO user_information
                    (esid, uid, service_address, unit_number, city, zip, status)
                VALUES (:esid, :uid, :address, :unit, :city, :zip, 0)
            """),
            {
                "esid": esid,
                "uid": uid,
                "address": address,
                "unit": unit_number,
                "city": city,
                "zip": zip_code,
            },
        )
        count += 1

    await db.commit()
    return count


# ---------------------------------------------------------------------------
# Admin — logs
# ---------------------------------------------------------------------------

async def get_all_logs(db: AsyncSession) -> list:
    result = await db.execute(
        text("""
            SELECT u.name AS customer_name,
                   ui.sr, ui.esid, ui.service_address, ui.unit_number,
                   ui.city, ui.zip, ui.status, ui.date_time
            FROM users u
            LEFT JOIN user_information ui ON u.uid = ui.uid
            WHERE ui.status IN (1, 2)
              AND ui.date_time IS NOT NULL
            ORDER BY ui.date_time DESC
        """)
    )
    rows = result.fetchall()
    status_map = {1: "Add", 2: "Cancel"}
    return [
        {
            "customer_name": r.customer_name,
            "sr": r.sr,
            "esid": r.esid,
            "service_address": r.service_address,
            "unit_number": r.unit_number,
            "city": r.city,
            "zip": r.zip,
            "status": status_map.get(r.status, "Unknown"),
            "date_time": str(r.date_time) if r.date_time else "",
        }
        for r in rows
    ]


# ---------------------------------------------------------------------------
# Customer — meters
# ---------------------------------------------------------------------------

async def list_meters(db: AsyncSession, uid: int) -> list:
    result = await db.execute(
        text("""
            SELECT sr, esid, service_address, unit_number, city, zip, status
            FROM user_information
            WHERE uid = :uid AND status != 4
            ORDER BY sr ASC
        """),
        {"uid": uid},
    )
    rows = result.fetchall()
    status_map = {0: "Pending", 1: "Add Requested", 2: "Cancel Requested", 3: "Failed"}
    return [
        {
            "sr": r.sr,
            "esid": r.esid,
            "service_address": r.service_address,
            "unit_number": r.unit_number,
            "city": r.city,
            "zip": r.zip,
            "status_code": r.status,
            "status": status_map.get(r.status, "Unknown"),
        }
        for r in rows
    ]


async def add_esiid(
    db: AsyncSession,
    uid: int,
    admin_name: str,
    esid: str,
    service_address: str,
    unit_number: str,
    city: str,
    zip_code: str,
) -> None:
    await db.execute(
        text("""
            INSERT INTO user_information
                (esid, uid, service_address, unit_number, city, zip, status)
            VALUES (:esid, :uid, :address, :unit, :city, :zip, 0)
        """),
        {
            "esid": esid,
            "uid": uid,
            "address": service_address,
            "unit": unit_number,
            "city": city,
            "zip": zip_code,
        },
    )

    today = date.today()
    await db.execute(
        text("""
            INSERT INTO user_log_add_meter
                (esid, uid, address, unit_number, city, zip, name, date, time)
            VALUES (:esid, :uid, :address, :unit, :city, :zip, :name, :dt, :tm)
        """),
        {
            "esid": esid,
            "uid": uid,
            "address": service_address,
            "unit": unit_number,
            "city": city,
            "zip": zip_code,
            "name": admin_name,
            "dt": today.strftime("%Y-%m-%d"),
            "tm": today.strftime("%H:%M:%S"),
        },
    )
    await db.commit()

    # Fetch customer name for email
    result = await db.execute(
        text("SELECT name FROM users WHERE uid = :uid"), {"uid": uid}
    )
    row = result.fetchone()
    customer_name = row.name if row else f"UID {uid}"

    _send_email(
        subject=f"New ESI ID Added — {customer_name}",
        html_body=f"""
        <p>A new location has been added to the Meter Portal.</p>
        <table>
          <tr><td><b>Customer:</b></td><td>{customer_name}</td></tr>
          <tr><td><b>ESI ID:</b></td><td>{esid}</td></tr>
          <tr><td><b>Address:</b></td><td>{service_address}</td></tr>
          <tr><td><b>Unit:</b></td><td>{unit_number}</td></tr>
          <tr><td><b>City:</b></td><td>{city}</td></tr>
          <tr><td><b>ZIP:</b></td><td>{zip_code}</td></tr>
          <tr><td><b>Added by:</b></td><td>{admin_name}</td></tr>
          <tr><td><b>Date:</b></td><td>{today.strftime("%m/%d/%Y")}</td></tr>
        </table>
        """,
    )


async def submit_request(
    db: AsyncSession,
    uid: int,
    srs: list,
    action: str,
    timing: str,
    custom_date: Optional[str],
    contact_name: str,
    contact_phone: str,
    contact_email: str,
    comments: str,
    customer_name: str,
) -> None:
    if not srs:
        raise HTTPException(status_code=400, detail="No meters selected")

    # Calculate effective date
    if timing == "same_day":
        effective_date = date.today().strftime("%m/%d/%Y")
    elif timing == "first_available":
        days = 3 if action == "cancel" else 2
        effective_date = _business_days_from_today(days)
    else:
        effective_date = custom_date or date.today().strftime("%m/%d/%Y")

    new_status = 1 if action == "add" else 2

    # Fetch selected meters for email
    placeholders = ", ".join(f":sr{i}" for i in range(len(srs)))
    params = {f"sr{i}": sr for i, sr in enumerate(srs)}
    params["uid"] = uid
    result = await db.execute(
        text(f"""
            SELECT sr, esid, service_address, unit_number, city, zip
            FROM user_information
            WHERE uid = :uid AND sr IN ({placeholders})
        """),
        params,
    )
    meters = result.fetchall()

    if not meters:
        raise HTTPException(status_code=400, detail="No valid meters found")

    # Update statuses
    for m in meters:
        await db.execute(
            text("""
                UPDATE user_information
                SET status = :status, date_time = NOW()
                WHERE sr = :sr AND uid = :uid
            """),
            {"status": new_status, "sr": m.sr, "uid": uid},
        )
    await db.commit()

    # Build email
    action_label = "Adding New Meters" if action == "add" else "Cancelling Meters"
    rows_html = "".join(
        f"<tr><td>{m.esid}</td><td>{m.service_address}</td>"
        f"<td>{m.unit_number}</td><td>{m.city}</td><td>{m.zip}</td></tr>"
        for m in meters
    )

    _send_email(
        subject=f"Meter {action_label} Request — {customer_name}",
        html_body=f"""
        <h3>METER AMENDMENT REQUEST</h3>
        <table>
          <tr><td><b>Customer:</b></td><td>{customer_name}</td></tr>
          <tr><td><b>Request Type:</b></td><td>{action_label}</td></tr>
          <tr><td><b>Effective Date:</b></td><td>{effective_date}</td></tr>
          <tr><td><b>Contact Name:</b></td><td>{contact_name}</td></tr>
          <tr><td><b>Contact Phone:</b></td><td>{contact_phone}</td></tr>
          <tr><td><b>Contact Email:</b></td><td>{contact_email}</td></tr>
          <tr><td><b>Comments:</b></td><td>{comments}</td></tr>
        </table>
        <br>
        <h4>Selected Meters</h4>
        <table border="1" cellpadding="5" style="border-collapse:collapse">
          <thead>
            <tr>
              <th>ESI ID</th><th>Address</th><th>Unit</th>
              <th>City</th><th>ZIP</th>
            </tr>
          </thead>
          <tbody>{rows_html}</tbody>
        </table>
        """,
    )
